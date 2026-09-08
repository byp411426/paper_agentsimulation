from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "output"
ASSET_DIR = OUT_DIR / "assets"
DOCX_PATH = OUT_DIR / "DisasterSociety_论文中文审阅稿.docx"
OUT_DIR.mkdir(exist_ok=True)
ASSET_DIR.mkdir(exist_ok=True)

FONT_PATH = Path(
    "/System/Library/AssetsV2/com_apple_MobileAsset_Font7/"
    "3419f2a427639ad8c8e139149a287865a90fa17e.asset/"
    "AssetData/PingFang.ttc"
)

NAVY = "17365D"
BLUE = "2E74B5"
LIGHT_BLUE = "EAF2F8"
LIGHT_GREEN = "EAF6EC"
LIGHT_ORANGE = "FFF3E0"
LIGHT_GRAY = "F3F5F7"
MID_GRAY = "6B7280"
WHITE = "FFFFFF"
BLACK = "111111"


def set_run_font(run, size=10.5, bold=False, italic=False, color=BLACK,
                 east_asia="PingFang SC", latin="Calibri"):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MID_GRAY)


def add_para(doc, text="", *, bold_prefix=None, italic=False, align=None,
             after=7, before=0, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.30
    p.paragraph_format.keep_together = keep
    p.alignment = align or WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    size = {1: 16, 2: 13, 3: 11.5}[level]
    color = BLUE if level < 3 else NAVY
    set_run_font(r, size=size, bold=True, color=color)
    return p


def add_equation(doc, text, number=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    set_run_font(r, size=10.5, east_asia="STSong", latin="Cambria Math")
    if number:
        r2 = p.add_run(f"    ({number})")
        set_run_font(r2, size=10, latin="Cambria Math")
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    r = p.add_run(text)
    set_run_font(r, size=9, color=MID_GRAY)
    return p


def add_table(doc, headers, rows, widths_dxa, font_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(header))
        set_run_font(r, size=font_size, bold=True, color=NAVY)
    set_repeat_table_header(table.rows[0])
    for ridx, row_data in enumerate(rows):
        cells = table.add_row().cells
        if ridx % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, "F8FAFC")
        for idx, value in enumerate(row_data):
            cell = cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    set_table_geometry(table, widths_dxa)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    return table


def load_font(size, bold=False):
    index = 1 if bold else 0
    return ImageFont.truetype(str(FONT_PATH), size=size, index=index)


def draw_centered_text(draw, box, text, font, fill="#17365D", line_gap=8):
    x0, y0, x1, y1 = box
    lines = text.split("\n")
    heights = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        heights.append(bbox[3] - bbox[1])
    total_h = sum(heights) + line_gap * (len(lines) - 1)
    y = y0 + (y1 - y0 - total_h) / 2
    for line, h in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=font)
        width = bbox[2] - bbox[0]
        draw.text((x0 + (x1 - x0 - width) / 2, y), line, font=font, fill=fill)
        y += h + line_gap


def rounded_box(draw, box, fill, outline="#4B5563", radius=16, width=3):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw, p1, p2, fill="#27364A", width=4):
    draw.line([p1, p2], fill=fill, width=width)
    import math
    angle = math.atan2(p2[1] - p1[1], p2[0] - p1[0])
    length = 14
    for delta in (2.55, -2.55):
        pt = (p2[0] + length * math.cos(angle + delta),
              p2[1] + length * math.sin(angle + delta))
        draw.line([p2, pt], fill=fill, width=width)


def make_overview(path):
    img = Image.new("RGB", (1800, 720), "white")
    d = ImageDraw.Draw(img)
    font = load_font(30)
    small = load_font(25)
    boxes = {
        "survey": (50, 100, 320, 225),
        "event": (50, 390, 320, 515),
        "adapter1": (400, 100, 710, 225),
        "adapter2": (400, 390, 710, 515),
        "schema": (790, 100, 1080, 225),
        "stream": (790, 390, 1080, 515),
        "core": (1160, 180, 1480, 360),
        "cal": (1160, 450, 1480, 610),
        "output": (1545, 180, 1760, 360),
        "audit": (1545, 450, 1760, 610),
    }
    for key in ("survey", "event", "output", "audit"):
        rounded_box(d, boxes[key], "#FFF3E0")
    for key in ("adapter1", "adapter2", "core"):
        rounded_box(d, boxes[key], "#EAF6EC")
    for key in ("schema", "stream", "cal"):
        rounded_box(d, boxes[key], "#EAF2F8")
    labels = {
        "survey": "灾后调查\n（结果仅用于评估）",
        "event": "独立官方\n灾害时间线",
        "adapter1": "数据集适配器\n字段角色与泄漏门控",
        "adapter2": "灾害适配器\n带来源的因果证据",
        "schema": "统一家庭\n人口模式",
        "stream": "带时间戳的\n事件流",
        "core": "通用智能体核心\n持续状态 · PADM认知\n合法动作 · 统一时钟",
        "cal": "校准层\n行为倾向 · 时间质量\n分层生命周期抽样",
        "output": "生成的家庭\n事件历史",
        "audit": "总体指标\n时间与审计日志",
    }
    for key, label in labels.items():
        draw_centered_text(d, boxes[key], label, font if key not in ("output", "audit") else small)
    arrow(d, (320, 162), (400, 162))
    arrow(d, (320, 452), (400, 452))
    arrow(d, (710, 162), (790, 162))
    arrow(d, (710, 452), (790, 452))
    arrow(d, (1080, 162), (1160, 235))
    arrow(d, (1080, 452), (1160, 305))
    arrow(d, (1320, 450), (1320, 360))
    arrow(d, (1480, 270), (1545, 270))
    arrow(d, (1652, 360), (1652, 450))
    img.save(path, dpi=(200, 200))


def make_agent_loop(path):
    img = Image.new("RGB", (1800, 630), "white")
    d = ImageDraw.Draw(img)
    font = load_font(29)
    boxes = [
        ((40, 85, 300, 215), "#FFF3E0", "时刻 t 的\n已送达证据"),
        ((360, 85, 635, 215), "#EAF6EC", "有界感知\n时间 · 来源 · 范围"),
        ((695, 85, 970, 215), "#EAF6EC", "PADM评估\n威胁 · 可行性 · 信任"),
        ((1030, 70, 1325, 230), "#EAF6EC", "持续状态\n阶段 · 认知\n可选记忆与计划"),
        ((1390, 85, 1750, 215), "#EAF6EC", "结构化LLM策略\n动作与准备度"),
        ((1390, 360, 1750, 520), "#EAF2F8", "延期质量控制器\npᵢ · wₜ · uᵢ,ₛ"),
        ((1030, 375, 1325, 505), "#EAF2F8", "合法执行器\n状态转移"),
        ((600, 375, 955, 505), "#FFF3E0", "轨迹日志\n与下一时刻反馈"),
    ]
    for box, fill, text in boxes:
        rounded_box(d, box, fill)
        draw_centered_text(d, box, text, font)
    for p1, p2 in [
        ((300, 150), (360, 150)),
        ((635, 150), (695, 150)),
        ((970, 150), (1030, 150)),
        ((1325, 150), (1390, 150)),
        ((1570, 215), (1570, 360)),
        ((1390, 440), (1325, 440)),
        ((1030, 440), (955, 440)),
    ]:
        arrow(d, p1, p2)
    d.line([(780, 375), (780, 285), (1170, 285), (1170, 230)],
           fill="#64748B", width=3)
    arrow(d, (1170, 285), (1170, 230), fill="#64748B", width=3)
    d.text((805, 255), "屏障提交后反馈", font=load_font(23), fill="#64748B")
    img.save(path, dpi=(200, 200))


overview_path = ASSET_DIR / "framework_overview_zh.png"
agent_path = ASSET_DIR / "agent_loop_zh.png"
make_overview(overview_path)
make_agent_loop(agent_path)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.82)
section.bottom_margin = Inches(0.82)
section.left_margin = Inches(0.95)
section.right_margin = Inches(0.95)
section.header_distance = Inches(0.42)
section.footer_distance = Inches(0.42)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.30

for level, size, before, after, color in [
    (1, 16, 16, 8, BLUE),
    (2, 13, 12, 6, BLUE),
    (3, 11.5, 8, 4, NAVY),
]:
    style = styles[f"Heading {level}"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hr = hp.add_run("DisasterSociety 论文中文审阅稿")
set_run_font(hr, size=8.5, color=MID_GRAY)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("中文内容审阅稿  ·  第 ")
set_run_font(fr, size=8.5, color=MID_GRAY)
add_page_field(fp)
fr2 = fp.add_run(" 页")
set_run_font(fr2, size=8.5, color=MID_GRAY)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(110)
p.paragraph_format.space_after = Pt(14)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("DisasterSociety")
set_run_font(r, size=30, bold=True, color=NAVY)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(20)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("面向跨灾种防护行为仿真的\n总体校准持续智能体框架")
set_run_font(r, size=20, bold=True, color=BLUE)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(70)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("论文中文内容审阅稿")
set_run_font(r, size=12.5, color=MID_GRAY)

for text in [
    "对应英文稿：IEEE Transactions / Journals 通用期刊格式",
    "研究对象：飓风与野火条件下的家庭防护行为事件历史",
    "当前状态：已纳入全部已完成定量实验；人工盲评尚未写入正文",
    "日期：2026年7月",
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    rr = p.add_run(text)
    set_run_font(rr, size=10.5, color=MID_GRAY)

doc.add_page_break()

add_heading(doc, "摘要", 1)
add_para(
    doc,
    "灾后调查能够记录家庭是否撤离以及何时撤离，但观测到的最终结果只是一个过程的终点。"
    "这个过程通常包括接收预警、判断信息来源、评估威胁、准备物资、协调家庭成员，以及最终执行或放弃防护行动。"
    "统计模型可以估计行为概率和总体分布，却不会自然生成上述连续过程；大语言模型智能体能够生成细致的行动轨迹，"
    "但其总体行为频率往往缺乏校准。本文提出 DisasterSociety：一种将“总体校准”与“过程生成”分离的生成式多智能体框架。"
    "框架首先用训练集拟合家庭连续行为倾向，再由持续存在的家庭智能体根据有时间边界且保留来源信息的证据更新"
    "防护行动决策模型（Protective Action Decision Model，PADM）状态。训练集拟合的时间质量只有在智能体进入合法的"
    "撤离前风险集合后才会被消耗；过早到达的时间质量会被延期，而不会被隐式转化为一个已发生事件。"
    "框架使用五分层、家庭偏移的拉丁超立方生命周期抽样降低有限种子带来的波动，同时不预先给家庭指定二元撤离类别。"
)
add_para(
    doc,
    "我们在 Hurricane Irma 与 Carr Fire 两个灾后调查数据集上进行验证，并使用独立公开事件时间线限制智能体可获得的信息。"
    "在 Irma 的 87 户测试队列中，五种子集成的 Macro-F1 为 0.551、总体撤离比例误差为 0.018；上下文倾向模型分别为"
    " 0.494 和 0.045。Macro-F1 点估计有所提高，但 95% 自助法置信区间跨越零。"
    "在配对的 30 户机制消融中，持续状态相对无状态智能体将 Macro-F1 提高 0.331，95% 置信区间为 [0.125, 0.522]；"
    "执行反馈、记忆和计划修订没有表现出可识别的额外预测收益。冻结的 Irma 到 Carr 迁移保留了部分家庭排序能力，"
    "但明显低估 Carr 的总体撤离基准率。仅使用 Carr 训练集拟合截距与时间适配器后，Carr 测试集总体比例误差降至"
    " 0.0029，Brier 分数降至 0.103，六小时时间分辨率下的撤离时间 Wasserstein 距离降至 5.49 小时。"
    "所有正式主运行均未出现非法状态转移或未来信息泄漏。结果支持持续状态以及“总体校准—过程生成分离”这两个机制，"
    "同时说明事件基准率和时间分布仍需要显式的灾害与事件适配器。"
)
add_para(doc, "关键词：生成式智能体；多智能体仿真；灾害撤离；防护行动；认知建模；事件历史；概率校准；跨灾种验证。",
         bold_prefix="关键词：")

add_heading(doc, "1 引言", 1)
add_para(
    doc,
    "家庭在野火或飓风中的防护行动不是一个瞬时分类问题。家庭会陆续接触到不完整的信息，判断信息来源是否可信，"
    "评估危险是否会影响自己，检查交通工具、照护对象和物资等约束，进行准备，随后才可能撤离或决定留下。"
    "单一的撤离标签无法表达这一历史；单独的撤离时间也无法解释家庭为何在该时刻具备行动条件。"
)
add_para(
    doc,
    "统计选择模型与事件历史模型能够提供经过校准的撤离概率和时间估计。当任务是监督预测时，它们仍然是必要的参考。"
    "但这类模型通常不会生成“观察—评估—准备—执行—反馈—修订”的因果序列。规则型智能体模型能够加入时间状态和交互"
    "[1], [2]，却要求研究者手工指定大量行为转移，并且往往与单一灾害绑定。"
)
add_para(
    doc,
    "大语言模型提供了另一种过程建模方式：受约束的 LLM 策略可以解释异构证据，并输出结构化行动。"
    "但是，文本看起来合理并不意味着总体分布正确。生成式智能体可能在语言层面高度可信，同时系统性地高估撤离、"
    "过早采取行动，或者在不同随机种子之间产生过大波动。如果研究目标是社会仿真，仅有轨迹可读性远远不够。"
)
add_para(
    doc,
    "本文把灾害防护行动仿真表述为两个相互关联但不可混同的问题。第一，总体校准问题：模拟群体的行为频率和时间分布"
    "应与调查数据一致。第二，过程生成问题：每个家庭必须按照当时可获得的信息、持续认知状态和现实约束，生成合法且可审计的"
    "事件历史。DisasterSociety 用统计层约束总体概率，用持续智能体生成过程，并用合法执行器连接二者。"
)
add_para(
    doc,
    "本文的主要贡献包括：提出一个跨灾种的统一家庭智能体接口，明确区分世界事实、智能体观察、智能体信念和评估标签；"
    "提出延期时间质量控制器，在不预先指定最终类别、不强制窗口末撤离的情况下，使连续家庭倾向与时间过程相结合；"
    "构建同步并发仿真内核，使智能体持续存在，同时避免并发带来的模拟时间错误；"
    "并在 Irma 与 Carr 数据上分别验证总体重建、持续状态机制、跨灾种迁移、目标训练适配与撤离时间重建。"
)

add_heading(doc, "2 相关工作", 1)
add_heading(doc, "2.1 灾害防护行为与撤离建模", 2)
add_para(
    doc,
    "PADM 将防护行动解释为由环境线索、社会线索、信息源、威胁感知、防护行动感知和利益相关者感知共同驱动的过程[7]。"
    "这一理论适合本研究，因为它明确区分“收到信息”“相信信息”“认为行动可行”和“真正执行行动”。"
    "传统离散选择模型适合估计是否撤离，生存分析或风险模型适合估计何时撤离[8]。"
    "真实野火撤离研究还表明，行为、道路网络和交通过程彼此关联[9]。相关综述指出，野火撤离仿真长期面临行为数据、"
    "决策模型与交通模型难以同时验证的问题[10]。本文不声称用 LLM 取代这些模型，而是研究持续生成过程如何与经验证的"
    "总体参考分布对接。"
)

add_heading(doc, "2.2 生成式智能体与社会仿真", 2)
add_para(
    doc,
    "Generative Agents 通过记忆、反思和规划生成长时段社会行为[3]。GenSim 将 LLM 智能体组织成通用社会实验平台[4]；"
    "AgentSociety 进一步关注大规模人口与干预实验[5]。已有综述指出，LLM 驱动的智能体建模具有更强的语义解释与交互能力，"
    "但校准、可重复性、计算成本和有效性仍是开放问题[11]。本研究借鉴持续状态与结构化认知，但不把“拥有更多模块”本身"
    "当作创新证据，而是通过消融检验各模块是否产生可识别影响。"
)

add_heading(doc, "2.3 LLM 与灾害行为", 2)
add_para(
    doc,
    "近期工作已经使用行为理论增强的 LLM 预测野火撤离决策[6]。这类研究的主要目标仍是个体分类。"
    "DisasterSociety 的目标不同：输出是整个群体在统一灾害时间线上的事件历史，而不是一次性标签。"
    "因此，个体分类只作为诊断指标之一，系统还必须评估总体比例、概率校准、撤离时间分布、状态转移合法性和信息因果性。"
)

add_heading(doc, "3 问题定义与系统边界", 1)
add_para(
    doc,
    "设模拟窗口包含离散时刻 t = 0, …, T−1，家庭集合为 i = 1, …, N。每个家庭包含人口属性 Xᵢ、"
    "经过合法信息传播后形成的观察 Oᵢ,ₜ、持续认知状态 Bᵢ,ₜ，以及动作 Aᵢ,ₜ。"
)
add_equation(doc, "ζᵢ,ₜ = (Xᵢ, Oᵢ,≤ₜ, Bᵢ,ₜ)", 1)
add_para(
    doc,
    "每个事件 e 都具有发布时间、地理范围、事件类型和来源。事件发布后并不等于所有家庭都已观察到它。"
    "世界状态 Wₜ、家庭观察 Oᵢ,ₜ 与家庭信念 Bᵢ,ₜ 必须保持区分："
)
add_equation(doc, "Wₜ ≠ Oᵢ,ₜ ≠ Bᵢ,ₜ", 2)
add_para(
    doc,
    "只有在事件已经于时刻 t 前公开、与家庭相关并通过预设投递规则时，事件才可进入观察。"
    "智能体策略不能看到调查结果、真实撤离时间、拟合行为倾向、随机抽样值或未来事件。"
)
add_para(
    doc,
    "经验目标由三部分构成：结果分布误差、时间分布误差和过程诊断。"
)
add_equation(doc, "L_outcome = d(P̂(Y), P_survey(Y))", 3)
add_equation(doc, "L_time = d(P̂(τ), P_survey(τ))", 4)
add_equation(doc, "L_process = φ({Hᵢ})", 5)
add_para(
    doc,
    "其中 d 表示分布距离或校准误差，φ 包含状态转移合法性、时间因果性与机制诊断。"
    "任何一个分类指标都不能单独替代这三类证据。"
)

add_heading(doc, "4 DisasterSociety 方法框架", 1)
add_heading(doc, "4.1 系统边界与适配器", 2)
add_para(
    doc,
    "数据集适配器将不同灾害调查中的变量映射为统一的家庭与结果模式；灾害适配器将官方命令和公共灾害记录映射为"
    "带类型、时间戳和来源的证据。通用仿真器只读取统一模式。结果字段保存在独立评估存储中，不能进入智能体观察。"
)
doc.add_picture(str(overview_path), width=Inches(6.45))
add_caption(doc, "图1  系统边界：调查结果与智能体输入分离；统一核心通过数据集和灾害适配器服务不同事件。")

add_heading(doc, "4.2 训练集拟合的家庭行为倾向", 2)
add_para(
    doc,
    "在源事件训练划分上拟合带正则化的上下文逻辑回归模型："
)
add_equation(doc, "pᵢ = P(Yᵢ=1 | Xᵢ) = σ(β₀ + βᵀXᵢ)", 6)
add_para(
    doc,
    "特征覆盖家庭构成、住房区域、交通能力、信息获得方式、年龄和决策角色等共同字段。"
    "缺失值仅使用训练划分估计的中位数填补。所得 pᵢ 是连续行为倾向，不是预先指定的撤离或留下标签。"
    "对于目标事件适配，仅允许在目标训练集上拟合一个截距 α："
)
add_equation(doc, "p′ᵢ = σ(logit(pᵢ) + α)", 7)
add_para(
    doc,
    "Carr 测试划分不会用于估计 α。该适配器仅修正事件总体基准率；家庭差异仍来自源事件中拟合的斜率。"
)

add_heading(doc, "4.3 训练集拟合的时间质量", 2)
add_para(
    doc,
    "令 nₜ 为训练集中在时间箱 t 内观察到的撤离数，采用 Jeffreys 平滑得到归一化时间权重："
)
add_equation(doc, "wₜ = (nₜ + 1/2) / Σₖ(nₖ + 1/2)，且 Σₜwₜ = 1", 8)
add_para(
    doc,
    "Irma 的主实验使用 12 小时时间步长；Carr 同时评估 12 小时和 6 小时分辨率。"
    "时间权重只表示该事件中撤离质量的相对分布，不表示家庭在该时刻必然撤离。"
)

add_heading(doc, "4.4 延期时间质量控制器", 2)
add_para(
    doc,
    "只有处于“准备中”或“协调中”的家庭才进入合法撤离前风险集合。令 Lᵢ,ₜ 表示家庭 i 在时刻 t 是否位于该集合，"
    "dᵢ,ₜ₋₁ 表示此前因尚未具备合法条件而延期的质量，cᵢ,ₜ₋₁ 表示已经消耗的累计质量。"
)
add_equation(doc, "qᵢ,ₜ = { wₜ + dᵢ,ₜ₋₁,  若 Lᵢ,ₜ=1；  0, 若 Lᵢ,ₜ=0 }", 9)
add_equation(doc, "dᵢ,ₜ = { 0, 若 Lᵢ,ₜ=1； dᵢ,ₜ₋₁+wₜ, 若 Lᵢ,ₜ=0 }", 10)
add_equation(doc, "cᵢ,ₜ = min(1, cᵢ,ₜ₋₁ + qᵢ,ₜ)", 11)
add_para(
    doc,
    "由此得到用于审计的离散风险："
)
add_equation(doc, "hᵢ,ₜ = pᵢqᵢ,ₜ / (1 − pᵢcᵢ,ₜ₋₁)", 12)
add_para(
    doc,
    "该设计解决了一个关键矛盾：如果家庭尚未完成准备，训练分布中过早到达的撤离质量不会消失，也不会暗中强迫家庭撤离，"
    "而是在其进入合法风险集合后继续被消耗。窗口结束时系统不会强制剩余家庭撤离。"
)

add_heading(doc, "4.5 降低方差的生命周期事件实现", 2)
add_para(
    doc,
    "每个家庭—种子对仅获得一个生命周期抽样值。正式实验使用 K=5 个分层种子，并根据稳定家庭哈希计算偏移 oᵢ∈{0,…,K−1}，"
    "再为家庭生成固定抖动 jᵢ∈[0,1)。"
)
add_equation(doc, "kᵢ,ₛ = (s + oᵢ) mod K", 13)
add_equation(doc, "uᵢ,ₛ = (kᵢ,ₛ + jᵢ) / K", 14)
add_para(
    doc,
    "家庭在首个满足 uᵢ,ₛ < pᵢcᵢ,ₜ 且状态合法的时刻撤离。"
    "这一五分层设计与拉丁超立方抽样思想一致[12]，在只运行少量种子时减少 Monte Carlo 波动，同时保留随机过程。"
    "同一家庭的种子被刻意分布到五个分层中，便于配对比较。"
)

add_heading(doc, "4.6 持续智能体状态与 PADM 评估", 2)
add_para(
    doc,
    "每个智能体跨时刻保留以下状态：当前阶段、威胁判断、保护行动判断、信息源信任、家庭能力、当前目标、"
    "上一动作的执行结果、可选的结构化记忆以及可选的短期计划。主评估配置持续保留合法阶段与 PADM 认知状态；"
    "机制消融再分别加入执行反馈、结构化记忆和计划修订。"
)
add_para(
    doc,
    "结构化记忆把事件转换为带时间戳的条目，只允许检索当前时刻之前的内容，并保存来源、可信度、地理范围和相关性。"
    "记忆不是任意对话历史，而是一个因果受限的证据存储。计划状态记录当前目标、依赖、最后修订原因和完成情况。"
    "执行器的接受、拒绝或完成结果只有在下一时刻才能反馈给策略。"
)

add_heading(doc, "4.7 合法动作状态机与执行器", 2)
add_para(
    doc,
    "家庭阶段遵循有限状态过程：未感知 → 监测 → 准备 → 协调 → 撤离；在有限窗口内也可以进入留下状态。"
    "每个阶段只有一组合法动作，例如收集信息、准备车辆、协调照护对象、声明准备完成或继续等待。"
    "LLM 只能提出动作和准备度，不能绕过控制器强制撤离。执行器检查动作与阶段是否兼容，提交合法转移，"
    "并记录接受、拒绝、目标更新和完成结果。"
)

add_heading(doc, "4.8 同步并发仿真内核", 2)
add_para(
    doc,
    "所有家庭共享同一个模拟时钟。每个时刻依次完成：回放外部事件并投递到期消息；冻结只读世界快照；"
    "构造有界观察；并发执行家庭决策；在任何决策影响其他家庭之前等待同步屏障；"
    "应用撤离控制器并验证动作；提交状态转移并排队后续消息；写入因果轨迹与有效性审计。"
    "正式实验最多并发发起 32 个模型调用。并发只改变真实运行耗时，不改变模拟时间。"
)
doc.add_picture(str(agent_path), width=Inches(6.45))
add_caption(doc, "图2  家庭智能体决策—执行闭环：智能体控制认知、准备和准备度；校准层控制事件实现；执行器保证状态合法。")

add_heading(doc, "4.9 审计与有效性门控", 2)
add_para(
    doc,
    "每次状态转移保存观察、认知前后状态、检索记忆、提议动作、受控动作、执行结果和撤离控制记录。"
    "每次模型调用保存提示哈希、随机种子、状态、令牌计数、延迟和缓存状态。输出必须通过结构化模式验证。"
    "格式错误允许一次修复；请求失败则使用规则策略并明确记录。若回退率超过 1%、存在缺失输出、未来事件进入观察，"
    "或非法状态转移超过预注册容忍度，该运行被判为无效。"
)

add_heading(doc, "5 实验设计", 1)
add_heading(doc, "5.1 数据与评估队列", 2)
add_para(
    doc,
    "实验包含两个真实灾害事件包。Irma 调查含 1,263 份有效回答和 921 份完成问卷。"
    "正式灾害内测试使用与 HEvOD 事件 21 对应的 Brevard County 公共信号队列，共 87 户，其中 42 户撤离、45 户留下。"
    "Carr Fire 调查含 647 份有效回答和 335 份完成问卷。正式 Carr 测试队列共 68 户，其中 60 户撤离、8 户留下。"
    "Carr 官方事件时间线来自 CAL FIRE。"
)
add_table(
    doc,
    ["字段", "Hurricane Irma", "Carr Fire"],
    [
        ["有效原始回答", "1,263", "647"],
        ["完成问卷", "921", "335"],
        ["有效撤离/留下结果", "917", "330"],
        ["完整撤离时间", "539", "292"],
        ["正式测试家庭", "87", "68"],
        ["测试集撤离", "42", "60"],
        ["测试集留下", "45", "8"],
        ["外部事件来源", "HEvOD", "CAL FIRE"],
    ],
    [3500, 2930, 2930],
)
add_caption(doc, "表1  数据集与评估队列")
add_para(
    doc,
    "两个调查均为便利样本，因此本文报告的是“以调查样本为条件的行为重建”，而不是县级或州级人口估计。"
    "数据不包含精确受访者位置、真实社会关系、GPS 路径和道路链路级交通测量。"
)

add_heading(doc, "5.2 数据划分、泄漏控制与方法冻结", 2)
add_para(
    doc,
    "每个事件包使用随机种子 42 进行家庭分层划分；同一受访者衍生的记录必须处于同一划分。"
    "调查结果、撤离时间和决策后字段被标记为目标或诊断，不能进入策略提示。"
    "Irma 训练划分用于拟合上下文倾向与时间权重；控制器和持续状态配置在生成 Irma 测试摘要之前冻结。"
)
add_para(
    doc,
    "由于项目早期曾查看 Carr 数据，本文不把 Carr 称为完全未接触的外部验证，而称为“方法冻结后的回顾性跨灾种确认”。"
    "冻结条件完全使用 Irma 的倾向与时间组件，不使用 Carr 标签。另设目标训练鲁棒性条件，只在 Carr 训练集上拟合截距"
    "和时间权重，绝不使用 Carr 测试标签。"
)

add_heading(doc, "5.3 基线与机制消融", 2)
add_para(
    doc,
    "Irma 灾害内比较包括五类方法：正则化逻辑回归结果模型、使用统一家庭特征的上下文倾向模型、"
    "非 LLM 的 PADM 规则智能体模型、一次性无状态 LLM，以及本文的总体校准持续过程。"
    "所有 LLM 实验使用 deepseek-v4-flash、温度 0.2、结构化 JSON 输出和注册种子 41—45。"
)
add_para(
    doc,
    "机制消融使用固定的 30 户 Irma 子集和四个累积配置：无状态、仅持续状态、持续状态加执行反馈、"
    "以及加入结构化记忆和计划修订的完整持续配置。每个配置使用五个配对种子。"
    "该设计检验具名模块是否真正改变经验结果，而不会把“系统里存在某模块”直接当作贡献。"
)

add_heading(doc, "5.4 指标与不确定性", 2)
add_para(
    doc,
    "结果指标包括 Macro-F1、平衡准确率、Brier 分数、期望校准误差（ECE）以及总体撤离比例绝对误差。"
    "五种子集成概率等于某家庭在五个种子中实现撤离的频率。"
    "撤离时间指标包括以小时计的 Wasserstein-1 距离、经验累积分布函数积分绝对误差（CDF IAE）和平均时间绝对误差。"
    "过程指标包括非法转移率、未来证据数、回退率、PADM 证据响应、来源评估覆盖率、准备—执行诊断、重复动作、"
    "反馈响应和计划修订。家庭层效应使用 10,000 次配对自助抽样；时间适配效应使用配对种子统计。"
    "若 95% 区间跨越零，则无论点估计方向如何，均判为结论不确定。"
)

add_heading(doc, "5.5 经验评估结构", 2)
add_table(
    doc,
    ["评估组成", "经验目的"],
    [
        ["数据有效性", "检查字段模式、来源、划分与信息泄漏"],
        ["Irma 灾害内重建", "评估结果分布与撤离时间"],
        ["机制消融", "检验持续状态、反馈、记忆与规划"],
        ["冻结跨灾种迁移", "进行 Irma 到 Carr 的回顾性确认"],
        ["目标训练鲁棒性", "评估 Carr 截距、时间与分辨率适配"],
        ["过程有效性", "检查状态序列合法性与信息因果性"],
    ],
    [2600, 6760],
)
add_caption(doc, "表2  论文中实际报告的评估组成")

add_heading(doc, "6 实验结果", 1)
add_heading(doc, "6.1 Irma 灾害内行为重建", 2)
add_table(
    doc,
    ["方法", "运行数", "Macro-F1", "平衡准确率", "Brier", "ECE", "比例误差"],
    [
        ["逻辑回归结果模型", "1", "0.490", "0.539", "0.248", "—", "0.096"],
        ["上下文倾向模型", "1", "0.494", "0.496", "0.261", "0.154", "0.045"],
        ["PADM 规则智能体", "5", "0.326", "0.500", "0.517", "—", "0.517"],
        ["一次性无状态 LLM", "5", "0.421", "0.507", "0.312", "—", "0.191"],
        ["DisasterSociety 校准持续过程", "5", "0.551", "0.553", "0.269", "0.124", "0.018"],
    ],
    [2740, 700, 1180, 1330, 1050, 950, 1410],
    font_size=8.0,
)
add_caption(doc, "表3  冻结的 87 户 Irma 测试队列结果")
add_para(
    doc,
    "总体校准持续过程取得最高 Macro-F1（0.551）、最高平衡准确率（0.553）和最低总体比例误差（0.018）。"
    "上下文倾向模型的 Brier 分数更低（0.261 对 0.269），而持续过程的 ECE 更低（0.124 对 0.154）。"
)
add_para(
    doc,
    "相对于上下文倾向模型，持续过程的 Macro-F1 效应为 +0.058，95% 家庭自助法区间为 [−0.023, 0.138]。"
    "点估计有利于持续过程，但该区间不能证明分类性能显著提高。倾向模型减持续过程的 Brier 效应为 −0.007，"
    "区间为 [−0.031, 0.015]，说明持续过程的 Brier 点估计略差。总体比例误差改善为 0.027，区间为 [−0.041, 0.047]。"
    "因此，持续过程改善了总体比例和 ECE 的点估计，但样本不足以支持全面统计优势。"
)
add_para(
    doc,
    "跨种子平均撤离比例为 0.501±0.052，真实调查比例为 0.483；种子级 Macro-F1 为 0.525±0.025。"
    "五个运行均有效，没有回退、模式修复、缺失输出、信息泄漏或非法状态转移。"
)

add_heading(doc, "6.2 持续状态是可识别的核心智能体机制", 2)
add_table(
    doc,
    ["配置", "Macro-F1", "平衡准确率", "Brier", "比例误差"],
    [
        ["无状态", "0.362", "0.500", "0.433", "0.433"],
        ["仅持续状态", "0.700", "0.717", "0.203", "0.053"],
        ["持续状态+执行反馈", "0.665", "0.688", "0.208", "0.080"],
        ["完整持续配置", "0.697", "0.699", "0.224", "0.027"],
    ],
    [2800, 1500, 1900, 1450, 1710],
)
add_caption(doc, "表4  五种子持续智能体机制消融")
add_para(
    doc,
    "无状态智能体不能跨时刻维持合法准备阶段，因此集成结果没有产生撤离，Macro-F1 仅为 0.362。"
    "仅加入持续状态后，Macro-F1 提高到 0.700，Brier 分数从 0.433 降到 0.203。"
    "仅持续状态相对无状态的家庭配对 Macro-F1 效应为 +0.331，95% 区间为 [0.125, 0.522]；"
    "Brier 改善为 +0.230，区间为 [0.039, 0.423]。两者均不跨零。"
)
add_para(
    doc,
    "加入执行反馈相对于仅持续状态的 Macro-F1 效应为 −0.035，区间为 [−0.116, 0.000]；"
    "完整持续配置相对于仅持续状态的效应为 −0.002，区间为 [−0.132, 0.130]。"
    "因此，现有数据支持“持续状态”这一机制，但没有证明反馈、记忆或计划修订带来额外预测收益。"
    "完整持续配置的总体比例误差点估计更低，但配对改善仍然不确定。"
)

add_heading(doc, "6.3 冻结的 Irma 到 Carr 跨灾种确认", 2)
add_table(
    doc,
    ["条件", "拟合数据", "Macro-F1", "平衡准确率", "Brier", "ECE", "比例误差"],
    [
        ["全部撤离", "无", "0.469", "0.500", "0.118", "0.118", "0.118"],
        ["Carr 上下文参考", "Carr训练", "0.469", "0.500", "0.116", "0.106", "0.021"],
        ["冻结 Irma 上下文模型", "Irma训练", "0.584", "0.679", "0.201", "0.321", "0.299"],
        ["冻结持续过程", "Irma训练", "0.595", "0.650", "0.197", "0.285", "0.285"],
        ["Carr截距+时间适配", "Irma+Carr训练", "0.469", "0.500", "0.103", "0.038", "0.0029"],
    ],
    [2390, 1390, 1120, 1350, 960, 880, 1270],
    font_size=7.8,
)
add_caption(doc, "表5  Carr 结果：冻结迁移与目标训练鲁棒性")
add_para(
    doc,
    "Carr 测试集真实撤离比例为 0.882，显著高于 Irma 队列的 0.483。冻结的 Irma 上下文倾向模型在 Carr 上仍保留"
    "一定家庭排序能力，Macro-F1 为 0.584、平衡准确率为 0.679，但预测总体比例只有 0.584。"
    "冻结持续过程的 Macro-F1 为 0.595、总体比例为 0.597。相对上下文倾向模型，其 Macro-F1 效应仅为 +0.012，"
    "95% 区间为 [−0.087, 0.092]；Brier 改善为 0.004，区间为 [−0.013, 0.022]，均不能得出显著优势。"
)
add_para(
    doc,
    "该结果区分了“排序能力”和“事件基准率校准”。冻结模型可以正确排序部分 Carr 家庭，但把总体比例低估了 0.285。"
    "只使用 Carr 训练集拟合截距后，过程集成比例提高到 0.885，与真实 0.882 的误差仅为 0.0029，Brier 降至 0.103。"
    "由于高基准率使几乎所有家庭概率都超过 0.5，Macro-F1 反而回落到多数类水平。"
    "这是一项总体校准修复，而不是个体分类性能提升。"
)

add_heading(doc, "6.4 撤离时间与适配器分辨率", 2)
add_table(
    doc,
    ["条件", "时间拟合", "步长", "Wasserstein/小时", "CDF IAE", "平均误差/小时"],
    [
        ["冻结迁移", "Irma训练", "12小时", "27.41", "0.2561", "25.10"],
        ["仅时间适配", "Carr训练", "12小时", "8.00", "0.0676", "3.43"],
        ["截距+时间适配", "Carr训练", "12小时", "7.40", "0.0628", "2.56"],
        ["截距+时间适配", "Carr训练", "6小时", "5.49", "0.0446", "2.73"],
    ],
    [2050, 1450, 1000, 1920, 1250, 1690],
    font_size=8.1,
)
add_caption(doc, "表6  Carr 撤离时间重建")
add_para(
    doc,
    "冻结的 Irma 时间分布迁移较差：Wasserstein 距离为 27.41 小时，CDF IAE 为 0.256。"
    "仅用 Carr 训练集替换时间权重后，Wasserstein 降到 8.00 小时，且不改变结果集成。"
    "联合截距与时间适配的 12 小时条件为 7.40 小时。保持拟合组件不变、把时间步长从 12 小时缩短到 6 小时后，"
    "Wasserstein 进一步降至 5.49 小时，CDF IAE 降至 0.0446。分辨率提升带来的配对 Wasserstein 改善为 1.91 小时，"
    "95% 种子区间为 [1.11, 2.69]。"
)
add_para(
    doc,
    "六小时条件把状态转移数从 3,060 增至 6,120，把记录的策略决策数从 2,376 增至 4,474；"
    "其结果指标与 12 小时联合适配条件完全一致。训练集交叉验证没有选择高斯平滑，因此平滑不作为独立实验条件。"
    "12 小时配置仍是默认的总体校准过程；6 小时配置只用于高分辨率时间分析。"
)

add_heading(doc, "6.5 序列合法性与过程诊断", 2)
add_para(
    doc,
    "Irma 校准过程的五次主运行均无非法执行。约 68.4% 的逐时刻控制器评估发生在合法撤离前风险集合之外，"
    "对应的时间质量因此被延期。在已经进入风险集合且智能体声明准备完成的决策中，平均“风险尚未触发”比例为 0.891。"
    "这说明准备度与事件实际实现是两个不同量。主 Irma 配置的 PADM 证据响应率和来源评估覆盖率均为 1.0，重复动作率为 0.412。"
)
add_para(
    doc,
    "Irma 撤离时间重建仍然偏弱。42 个具有时间标签的真实撤离中，有 10 个发生在选定公共信号之前，"
    "因此不参与信号后的时间归因。对其余撤离，12 小时过程的 Wasserstein 距离为 17.69 小时，CDF IAE 为 0.231，"
    "平均时间误差为 16.52 小时。该轨迹在程序上可以合法执行，但当前公共信号和粗粒度时间模型不能准确重建 Irma 的撤离时刻。"
)
add_para(
    doc,
    "五个冻结 Carr 迁移运行包含 3,060 次状态转移，没有时间泄漏、动作拒绝或回退。"
    "12 小时联合适配同样包含 3,060 次合法转移；6 小时条件包含 6,120 次合法转移。"
    "这些检查证明的是仿真器有效性，而不是行为真实性。"
)

add_heading(doc, "7 讨论", 1)
add_heading(doc, "7.1 现有结果真正支持什么", 2)
add_para(
    doc,
    "最清晰的机制证据来自持续状态消融。无状态策略不能跨时刻维持准备过程，也无法稳定进入合法撤离风险集合。"
    "持续保存阶段和 PADM 认知后，Macro-F1 提高 0.331，且置信区间不跨零。"
    "这一效应证明持续性在当前仿真器中的价值，但不能被扩张为对所有 LLM 智能体任务的普遍结论。"
)
add_para(
    doc,
    "Irma 主实验给出了另一类结果：持续过程在总体比例和 ECE 点估计上更接近调查数据，但 Brier 略差，"
    "Macro-F1 置信区间跨零。总体校准与个体分类不能互相替代。一个过程模型可能生成更有用的群体实现，"
    "但并不因此成为更好的监督分类器。"
)
add_para(
    doc,
    "Carr 实验揭示了跨灾种迁移的边界。Irma 训练的家庭斜率保留了一部分区分能力，但 Carr 具有完全不同的总体基准率"
    "和撤离时间分布。一个仅用目标训练集拟合的截距和时间直方图就能修复这些量，而无需改变通用智能体状态和动作契约。"
    "因此，可复用对象不是“一套对所有灾害都固定的概率分布”，而是“共享认知—执行核心加显式灾害与事件适配器”。"
)

add_heading(doc, "7.2 为什么更多认知模块没有提高预测", 2)
add_para(
    doc,
    "完整持续配置包含执行反馈、结构化记忆和计划修订，但并未优于仅持续状态。可能原因包括："
    "Irma 的 72 小时短窗口不需要长期记忆检索；可用公共信号本身稀疏，丰富记忆无法创造不存在的证据；"
    "总体控制器已经决定了有限窗口内的大部分结果分布，使认知模块更可能改变序列质量而不是最终标签。"
    "当前定量数据无法识别这种序列质量收益。"
)
add_para(
    doc,
    "因此，本文不把记忆、反馈和计划修订列为已经独立验证的贡献。它们仍是具有明确输入、输出和诊断接口的可选机制。"
    "这一负结果具有价值，因为它阻止我们把架构复杂度错误地包装成经验证的科学贡献。"
)

add_heading(doc, "7.3 对生成式社会仿真的启示", 2)
add_para(
    doc,
    "本框架把 LLM 定位为受约束的过程策略：LLM 负责解释证据和提出合法行动，但不能控制环境事实、总体基准率和随机事件实现。"
    "这一边界回应了生成式社会仿真中的常见验证失败：个体文本可以非常可信，而总体分布仍然完全错误。"
)
add_para(
    doc,
    "延期质量接口还使模型之间的不一致保持可观测。如果智能体始终未准备，其统计倾向不会在窗口结束时暗中强迫其撤离；"
    "如果智能体已经准备但生命周期抽样值仍高于累计消耗质量，准备完成也不会强制事件发生。"
    "这些情况都会保留在轨迹日志中，可作为过程摩擦诊断。"
)

add_heading(doc, "8 局限性与伦理边界", 1)
add_para(
    doc,
    "两个调查均为便利样本并依赖自报行为。正式队列规模较小，尤其是 30 户机制消融和 Carr 测试集中仅 8 户的留下少数类。"
    "自助法区间只反映这些队列内部的抽样不确定性，不能把它们转化为具有总体代表性的样本。"
)
add_para(
    doc,
    "Carr 在项目早期开发中曾被查看，因此跨灾种结果只能表述为方法冻结后的回顾性确认，而不是完全前瞻、从未接触的外部验证。"
    "Carr 训练适配器使用目标训练标签，只能解释为鲁棒性分析，不能称为零样本泛化。"
)
add_para(
    doc,
    "事件时间线仍较粗糙。Irma 时间分析从一个 HEvOD 公共信号开始，但队列中 23.8% 的有时间标签撤离发生在该信号之前。"
    "Carr 只有在使用灾害特定时间适配器后才获得明显改善，这进一步证明撤离时钟具有事件依赖性。"
)
add_para(
    doc,
    "仿真器不能恢复真实社会关系；社会图消息传播属于可选机制，在正式校准配置中关闭。"
    "本文不模拟具有物理精度的野火蔓延、飓风风场、交通流、道路容量、最优路线或清空时间，"
    "因为现有调查没有受访者 GPS 路径或链路级交通真值。"
)
add_para(
    doc,
    "LLM 通过第三方转发服务调用，其不可变模型快照标识无法独立验证。项目保存提示、输出、种子、哈希与缓存记录，"
    "但仍无法完全重建服务提供方内部状态。调查变量被转换为有限行为角色；这些角色是工程接口和审计记录，"
    "并不代表对真实人类认知状态的直接测量。"
)
add_para(
    doc,
    "合成智能体结果不应直接用于发布撤离命令、分配应急资源或推断某个具名受访者的行为。"
    "系统用于明确不确定性下的方法研究与情景分析；任何实际运行解释仍应由人类决策者负责。"
)

add_heading(doc, "9 结论", 1)
add_para(
    doc,
    "DisasterSociety 将训练集拟合的总体模型、持续且证据有界的家庭智能体以及合法事件历史执行器结合起来。"
    "延期时间质量控制器在不预先给家庭指定撤离类别、不在窗口末强制撤离的情况下保留连续家庭倾向；"
    "家庭偏移的五分层生命周期抽样降低有限种子波动；同步时钟内核使未来信息泄漏和非法状态转移可被审计。"
)
add_para(
    doc,
    "Irma 与 Carr 实验支持持续状态作为当前最主要的智能体机制。Irma 的总体比例和 ECE 点估计有所改善，"
    "但尚未证明个体分类显著优于上下文倾向模型。冻结的 Irma 到 Carr 迁移保留部分区分能力，却无法迁移 Carr 的基准率"
    "和时间分布；少量 Carr 训练集适配能够修复总体校准与时间。结果表明，跨灾种仿真器需要共享过程核心，也需要显式的"
    "事件级参数。记忆、反馈与计划修订在当前数据中没有表现出可识别的额外预测价值。"
)

add_heading(doc, "参考文献", 1)
references = [
    '[1] E. Bonabeau, “Agent-based modeling: Methods and techniques for simulating human systems,” PNAS, 2002.',
    '[2] C. M. Macal and M. J. North, “Tutorial on agent-based modelling and simulation,” Journal of Simulation, 2010.',
    '[3] J. S. Park et al., “Generative agents: Interactive simulacra of human behavior,” UIST, 2023.',
    '[4] J. Tang et al., “GenSim: A general social simulation platform with large language model based agents,” arXiv:2410.04360, 2024.',
    '[5] J. Piao et al., “AgentSociety: Large-scale simulation of LLM-driven generative agents,” arXiv:2502.08691, 2025.',
    '[6] R. Chen et al., “From perceptions to decisions: Wildfire evacuation decision prediction with behavioral theory-informed LLMs,” arXiv:2502.17701, 2025.',
    '[7] M. K. Lindell and R. W. Perry, “The protective action decision model: Theoretical modifications and additional evidence,” Risk Analysis, 2012.',
    '[8] D. R. Cox, “Regression models and life-tables,” Journal of the Royal Statistical Society: Series B, 1972.',
    '[9] T. Toledo et al., “Analysis of evacuation behavior in a wildfire event,” International Journal of Disaster Risk Reduction, 2018.',
    '[10] S. M. Beyki et al., “Evacuation simulation under threat of wildfire: An overview of research, development, and knowledge gaps,” Applied Sciences, 2023.',
    '[11] C. Gao et al., “Large language models empowered agent-based modeling and simulation: A survey and perspectives,” arXiv:2312.11970, 2023.',
    '[12] M. D. McKay, R. J. Beckman, and W. J. Conover, “A comparison of three methods for selecting values of input variables,” Technometrics, 1979.',
]
for ref in references:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(ref)
    set_run_font(r, size=8.8)

add_heading(doc, "审阅说明", 1)
add_para(
    doc,
    "本文件用于中文内容审阅，不是最终投稿排版。正文只包含已经完成并能够由现有数据支撑的定量实验。"
    "内部开发版本号和内部实验流水号均未写入；尚未执行的人工盲评也未作为实验或结果写入。"
    "英文 IEEE 稿仍是最终投稿格式的主文件。"
)

doc.core_properties.title = "DisasterSociety：面向跨灾种防护行为仿真的总体校准持续智能体框架"
doc.core_properties.subject = "论文中文内容审阅稿"
doc.core_properties.author = "Anonymous Author(s)"
doc.core_properties.keywords = "生成式智能体, 灾害仿真, 多智能体系统, 防护行动, 概率校准"
doc.core_properties.comments = "中文审阅稿；不包含尚未完成的人工盲评。"

alt_texts = [
    "DisasterSociety系统边界图：灾后调查和独立灾害时间线分别经数据集适配器与灾害适配器进入统一家庭模式、事件流、通用智能体核心和校准层，最终输出家庭事件历史与审计指标。",
    "家庭智能体决策执行闭环图：已送达证据依次进入有界感知、PADM评估、持续状态和结构化LLM策略，再经延期质量控制器与合法执行器写入轨迹日志并在下一时刻反馈。",
]
for shape, alt_text in zip(doc.inline_shapes, alt_texts):
    shape._inline.docPr.set("title", alt_text.split("：", 1)[0])
    shape._inline.docPr.set("descr", alt_text)

doc.save(DOCX_PATH)
print(DOCX_PATH)
